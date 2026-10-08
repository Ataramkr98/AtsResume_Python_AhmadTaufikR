from __future__ import annotations

import io
import uuid

import pytest
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.urls import reverse
from docx import Document
from icalendar import Calendar

from apps.accounts.models import Invitation, Organization, OrganizationMembership, User
from apps.applications.models import Application, ApplicationStageHistory, RejectionReason
from apps.applications.services import move_application_stage, recompute_match_score
from apps.candidates.models import Candidate, DuplicateSuggestion
from apps.candidates.services import merge_candidates
from apps.interviews.models import Interview, InterviewParticipant, ScorecardSubmission
from apps.interviews.services import save_scorecard
from apps.interviews.tasks import send_interview_invites
from apps.jobs.models import (
    Department,
    Job,
    JobSkillRequirement,
    Location,
    PipelineStage,
    PipelineTemplate,
    ScorecardCriterion,
    ScorecardTemplate,
    Skill,
    StageType,
)
from apps.tasks.models import Task


@pytest.fixture
def organization():
    return Organization.objects.create(name="Acme", slug="acme", job_approval_required=False)


@pytest.fixture
def admin_user(organization):
    user = User.objects.create_user("admin@acme.test", "LongPassword123!", full_name="Admin User")
    membership = OrganizationMembership.objects.create(
        organization=organization, user=user, role=OrganizationMembership.Role.ORG_ADMIN
    )
    return user, membership


@pytest.fixture
def job(organization, admin_user):
    user, _ = admin_user
    department = Department.objects.create(organization=organization, name="Engineering")
    location = Location.objects.create(
        organization=organization,
        name="Jakarta",
        city="Jakarta",
        country_code="ID",
        workplace_type=Location.WorkplaceType.HYBRID,
    )
    template = PipelineTemplate.objects.create(organization=organization, name="Default")
    return Job.objects.create(
        organization=organization,
        department=department,
        location=location,
        hiring_manager=user,
        recruiter=user,
        pipeline_template=template,
        title="Backend Engineer",
        slug="backend-engineer",
        status=Job.Status.OPEN,
        employment_type=Job.EmploymentType.FULL_TIME,
        description_html="<p>Build services.</p>",
        min_experience_months=48,
        education_level="bachelor",
    )


@pytest.fixture
def stages(job):
    applied = PipelineStage.objects.create(
        job=job, name="Applied", stage_type=StageType.APPLIED, position=1
    )
    screening = PipelineStage.objects.create(
        job=job, name="Screening", stage_type=StageType.SCREENING, position=2
    )
    return applied, screening


@pytest.fixture
def candidate(organization):
    candidate = Candidate(
        organization=organization,
        first_name="Nadia",
        last_name="Putri",
        city="Jakarta",
        country_code="ID",
        total_experience_months=60,
        education_level="bachelor",
    )
    candidate.set_email("nadia@example.test")
    candidate.save()
    return candidate


@pytest.fixture
def application(organization, job, stages, candidate, admin_user):
    user, _ = admin_user
    return Application.objects.create(
        organization=organization,
        job=job,
        candidate=candidate,
        current_stage=stages[0],
        source="MANUAL",
        assigned_recruiter=user,
    )


@pytest.mark.django_db
def test_tenant_boundary_returns_404(client, organization, admin_user):
    user, _ = admin_user
    other = Organization.objects.create(name="Other", slug="other")
    client.force_login(user)
    response = client.get(reverse("dashboard", args=[other.slug]))
    assert response.status_code == 404


@pytest.mark.django_db
def test_admin_mfa_enforcement_redirects_to_enrollment(client, settings, organization, admin_user):
    settings.MFA_ENFORCEMENT = True
    user, _ = admin_user
    client.force_login(user)
    response = client.get(reverse("dashboard", args=[organization.slug]))
    assert response.status_code == 302
    assert response.url == reverse("mfa-setup")
    setup = client.get(response.url)
    assert setup.status_code == 200
    assert b"multi-factor authentication" in setup.content


@pytest.mark.django_db
def test_stage_move_is_versioned_idempotent_and_audited(
    organization, admin_user, application, stages
):
    user, membership = admin_user
    moved = move_application_stage(
        application_id=application.id,
        target_stage_id=stages[1].id,
        organization=organization,
        actor=user,
        membership=membership,
        idempotency_key="move-1",
        expected_version=1,
    )
    assert moved.current_stage == stages[1]
    assert moved.version == 2
    assert ApplicationStageHistory.objects.filter(application=application, to_stage=stages[1]).count() == 1
    repeated = move_application_stage(
        application_id=application.id,
        target_stage_id=stages[1].id,
        organization=organization,
        actor=user,
        membership=membership,
        idempotency_key="move-1",
        expected_version=1,
    )
    assert repeated.id == application.id
    assert ApplicationStageHistory.objects.filter(application=application, to_stage=stages[1]).count() == 1
    with pytest.raises(ValidationError, match="changed"):
        move_application_stage(
            application_id=application.id,
            target_stage_id=stages[0].id,
            organization=organization,
            actor=user,
            membership=membership,
            idempotency_key="move-2",
            expected_version=1,
        )
    with pytest.raises(ValidationError, match="already in this stage"):
        move_application_stage(
            application_id=application.id,
            target_stage_id=stages[1].id,
            organization=organization,
            actor=user,
            membership=membership,
            idempotency_key="move-3",
            expected_version=2,
        )


@pytest.mark.django_db
def test_explainable_matching_excludes_unconfigured_dimensions(organization, job, candidate, application):
    python = Skill.objects.create(organization=organization, name="Python")
    django = Skill.objects.create(organization=organization, name="Django")
    JobSkillRequirement.objects.create(
        job=job, skill=python, requirement_type=JobSkillRequirement.RequirementType.MUST
    )
    JobSkillRequirement.objects.create(
        job=job, skill=django, requirement_type=JobSkillRequirement.RequirementType.PREFERRED
    )
    candidate.candidate_skills.create(skill=python, source="MANUAL")
    score = recompute_match_score(application)
    assert 70 <= score.total_score <= 100
    assert score.explanation["matched"] == ["Python"]
    assert score.explanation["missing"] == ["Django"]
    assert "automatic hiring decision" in score.explanation["disclaimer"]


@pytest.mark.django_db
def test_scorecard_extreme_rating_requires_evidence_and_locks(admin_user, application, job):
    user, _ = admin_user
    template = ScorecardTemplate.objects.create(job=job, name="Default", version=1)
    criterion = ScorecardCriterion.objects.create(
        template=template, competency="Expertise", position=1
    )
    interview = Interview.objects.create(
        application=application,
        scorecard_template=template,
        template_snapshot={
            "name": "Default",
            "version": 1,
            "criteria": [{"id": str(criterion.id), "competency": "Expertise", "position": 1}],
        },
        title="Technical interview",
        start_at="2026-08-01T10:00:00Z",
        end_at="2026-08-01T11:00:00Z",
        created_by=user,
    )
    InterviewParticipant.objects.create(interview=interview, user=user)
    with pytest.raises(ValidationError, match="Evidence"):
        save_scorecard(
            interview=interview,
            actor=user,
            recommendation="STRONG_YES",
            summary="Excellent",
            answers={str(criterion.id): {"rating": 5, "evidence": ""}},
            submit=True,
        )
    submission = save_scorecard(
        interview=interview,
        actor=user,
        recommendation="STRONG_YES",
        summary="Excellent",
        answers={str(criterion.id): {"rating": 5, "evidence": "Designed a comparable system."}},
        submit=True,
    )
    assert submission.status == ScorecardSubmission.Status.SUBMITTED
    assert submission.locked_at is not None
    with pytest.raises(ValidationError, match="locked"):
        save_scorecard(
            interview=interview,
            actor=user,
            recommendation="YES",
            summary="Changed",
            answers={str(criterion.id): {"rating": 4, "evidence": ""}},
        )


@pytest.mark.django_db
def test_interview_invite_task_sends_a_calendar_invite(demo, mailoutbox):
    """The invite task builds a real iCalendar attachment per panellist.

    `schedule_interview` enqueues it through `transaction.on_commit`, and those callbacks never
    fire inside a transactional test, so this path was previously never executed at all. Calling
    the task directly is what actually exercises it.
    """
    interview = (
        Interview.objects.filter(application__organization=demo)
        .select_related("application__candidate", "application__job")
        .prefetch_related("participants__user")
        .first()
    )
    assert interview is not None, "the demo workspace should contain interviews"
    participants = list(interview.participants.select_related("user"))
    assert participants, "the demo interview should have panellists"

    interview.meeting_url = "https://meet.example.test/abc-defg-hij"
    interview.save(update_fields=["meeting_url", "updated_at"])

    send_interview_invites(str(interview.id))

    assert len(mailoutbox) == len(participants), "every panellist should be invited"
    message = mailoutbox[0]
    assert message.subject == f"Interview: {interview.title}"
    assert message.to == [participants[0].user.email]
    assert str(interview.application.candidate) in message.body
    assert interview.application.job.title in message.body

    assert len(message.attachments) == 1
    filename, content, mimetype = message.attachments[0]
    assert filename == "interview.ics"
    assert mimetype == "text/calendar"

    # Parse it back: an attachment that is not valid iCalendar is worse than none at all.
    events = list(Calendar.from_ical(content).walk("VEVENT"))
    assert len(events) == 1
    assert str(events[0]["SUMMARY"]) == interview.title
    assert str(events[0]["LOCATION"]) == interview.meeting_url


@pytest.mark.django_db(transaction=True)
def test_public_apply_is_generic_and_deduplicates(client, settings, tmp_path, organization, job, stages):
    settings.MEDIA_ROOT = tmp_path
    settings.CELERY_TASK_ALWAYS_EAGER = True
    document = Document()
    document.add_paragraph("Nadia Putri")
    document.add_paragraph("nadia@example.test")
    document.add_heading("Skills", level=1)
    document.add_paragraph("Python and Django")
    buffer = io.BytesIO()
    document.save(buffer)

    def upload():
        return SimpleUploadedFile(
            "resume.docx",
            buffer.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

    endpoint = reverse("public-apply", args=[organization.slug, job.slug])
    payload = {
        "first_name": "Nadia",
        "last_name": "Putri",
        "email": "nadia@example.test",
        "phone": "+62 812 5555 0101",
        "consent": "on",
    }
    first = client.post(endpoint, {**payload, "resume": upload()})
    second = client.post(endpoint, {**payload, "resume": upload()})
    assert first.status_code == 202
    assert b"Application received" in first.content
    assert second.status_code == 202
    assert Candidate.objects.filter(organization=organization).count() == 1
    assert Application.objects.filter(job=job).count() == 1


@pytest.mark.django_db
def test_seed_demo_is_idempotent(client):
    """``compose.yaml`` runs ``seed_demo`` on every container start, so a second run must
    be a no-op. This guards the state-machine walk, which used to try to reopen a CLOSED
    requisition and abort the whole command."""
    call_command("seed_demo", verbosity=0)
    organization = Organization.objects.get(slug="northstar-labs")

    def snapshot():
        return (
            Job.objects.filter(organization=organization).count(),
            Candidate.objects.filter(organization=organization).count(),
            Application.objects.filter(organization=organization).count(),
            Interview.objects.filter(application__organization=organization).count(),
            Task.objects.filter(organization=organization).count(),
        )

    before = snapshot()
    closed_before = sorted(
        Job.objects.filter(organization=organization, status=Job.Status.CLOSED).values_list(
            "slug", flat=True
        )
    )
    assert closed_before, "the demo data should include at least one closed requisition"

    call_command("seed_demo", verbosity=0)

    assert snapshot() == before
    assert sorted(
        Job.objects.filter(organization=organization, status=Job.Status.CLOSED).values_list(
            "slug", flat=True
        )
    ) == closed_before


@pytest.mark.django_db
def test_seeded_primary_pages_render(client):
    call_command("seed_demo", verbosity=0)
    assert client.login(username="admin@example.com", password="DemoPass123!")
    organization = Organization.objects.get(slug="northstar-labs")
    job = Job.objects.get(organization=organization, slug="senior-backend-engineer")
    candidate = Candidate.objects.filter(organization=organization).first()
    paths = [
        reverse("dashboard", args=[organization.slug]),
        reverse("jobs", args=[organization.slug]),
        reverse("job-create", args=[organization.slug]),
        reverse("job-detail", args=[organization.slug, job.id]),
        reverse("job-edit", args=[organization.slug, job.id]),
        reverse("candidates", args=[organization.slug]),
        reverse("candidate-create", args=[organization.slug]),
        reverse("candidate-detail", args=[organization.slug, candidate.id]),
        reverse("pipeline", args=[organization.slug]),
        reverse("interviews", args=[organization.slug]),
        reverse("tasks", args=[organization.slug]),
        reverse("task-create", args=[organization.slug]),
        reverse("reports", args=[organization.slug]),
        reverse("profile", args=[organization.slug]),
        reverse("organization-settings", args=[organization.slug]),
        reverse("audit-log", args=[organization.slug]),
        reverse("public-jobs", args=[organization.slug]),
        reverse("public-job-detail", args=[organization.slug, job.slug]),
    ]
    for path in paths:
        response = client.get(path)
        assert response.status_code == 200, path


@pytest.mark.django_db
def test_editorial_auth_and_interview_states_render(client):
    call_command("seed_demo", verbosity=0)
    organization = Organization.objects.get(slug="northstar-labs")
    job = Job.objects.get(organization=organization, slug="senior-backend-engineer")
    application = Application.objects.filter(organization=organization, job=job).first()
    assert client.get(reverse("login")).status_code == 200
    assert client.get(reverse("public-apply", args=[organization.slug, job.slug])).status_code == 200
    assert client.get(reverse("accept-invitation", args=["invalid-token"])).status_code == 410

    assert client.login(username="admin@example.com", password="DemoPass123!")
    assert client.get(reverse("mfa-setup")).status_code == 200
    assert client.get(
        reverse("interview-create", args=[organization.slug, application.id])
    ).status_code == 200

    template = job.scorecard_templates.first()
    interview = Interview.objects.create(
        application=application,
        scorecard_template=template,
        template_snapshot={
            "name": template.name,
            "version": template.version,
            "criteria": [
                {"id": str(criterion.id), "competency": criterion.competency, "position": criterion.position}
                for criterion in template.criteria.all()
            ],
        },
        title="Demo interview",
        start_at="2026-08-25T10:00:00Z",
        end_at="2026-08-25T11:00:00Z",
        created_by=User.objects.get(email="admin@example.com"),
    )
    InterviewParticipant.objects.create(
        interview=interview,
        user=User.objects.get(email="admin@example.com"),
    )
    scorecard_response = client.get(
        reverse("scorecard", args=[organization.slug, interview.id])
    )
    assert scorecard_response.status_code == 200
    assert b"Save draft" in scorecard_response.content

    _, token = Invitation.issue(
        organization=organization,
        email="new.hire@example.com",
        role=OrganizationMembership.Role.RECRUITER,
        invited_by=User.objects.get(email="admin@example.com"),
    )
    assert client.get(reverse("accept-invitation", args=[token])).status_code == 200


@pytest.mark.django_db
def test_interviewer_only_sees_assigned_candidate_across_web_and_api(
    client, organization, job, stages, application
):
    interviewer = User.objects.create_user(
        "interviewer@acme.test", "LongPassword123!", full_name="Assigned Interviewer"
    )
    OrganizationMembership.objects.create(
        organization=organization,
        user=interviewer,
        role=OrganizationMembership.Role.INTERVIEWER,
    )
    hidden_candidate = Candidate(organization=organization, first_name="Hidden", last_name="Person")
    hidden_candidate.set_email("hidden@example.test")
    hidden_candidate.save()
    Application.objects.create(
        organization=organization,
        job=job,
        candidate=hidden_candidate,
        current_stage=stages[0],
        source="MANUAL",
    )
    template = ScorecardTemplate.objects.create(job=job, name="Interview", version=1)
    interview = Interview.objects.create(
        application=application,
        scorecard_template=template,
        template_snapshot={"name": "Interview", "version": 1, "criteria": []},
        title="Assigned interview",
        start_at="2026-08-25T10:00:00Z",
        end_at="2026-08-25T11:00:00Z",
        created_by=job.hiring_manager,
    )
    InterviewParticipant.objects.create(interview=interview, user=interviewer)
    client.force_login(interviewer)

    candidates_page = client.get(reverse("candidates", args=[organization.slug]))
    assert candidates_page.status_code == 200
    assert b"Nadia Putri" in candidates_page.content
    assert b"Hidden Person" not in candidates_page.content
    pipeline_page = client.get(reverse("pipeline-job", args=[organization.slug, job.id]))
    assert b"Nadia Putri" in pipeline_page.content
    assert b"Hidden Person" not in pipeline_page.content
    assert b"Reject application" not in pipeline_page.content
    assert client.get(
        reverse("candidate-detail", args=[organization.slug, hidden_candidate.id])
    ).status_code == 404
    assert client.get(reverse("organization-settings", args=[organization.slug])).status_code == 404
    assert b"Create job" not in client.get(
        reverse("dashboard", args=[organization.slug])
    ).content

    api_candidates = client.get(reverse("api-candidates", args=[organization.slug])).json()
    api_applications = client.get(reverse("api-applications", args=[organization.slug])).json()
    assert api_candidates["success"] is True
    assert api_applications["success"] is True
    assert [row["id"] for row in api_candidates["data"]["results"]] == [str(application.candidate_id)]
    assert [row["id"] for row in api_applications["data"]["results"]] == [str(application.id)]


@pytest.mark.django_db
def test_public_routes_hide_suspended_and_deleted_jobs(client, organization, job):
    organization.status = Organization.Status.SUSPENDED
    organization.save(update_fields=["status", "updated_at"])
    assert client.get(reverse("public-job-detail", args=[organization.slug, job.slug])).status_code == 404
    suspended_apply = client.post(
        reverse("api-public-apply", args=[organization.slug, job.slug]),
        {
            "first_name": "Applicant",
            "last_name": "Person",
            "email": "applicant@example.test",
            "phone": "",
            "resume": SimpleUploadedFile("resume.docx", b"PKfake"),
            "consent": True,
        },
        HTTP_IDEMPOTENCY_KEY="suspended-job",
    )
    assert suspended_apply.status_code == 404
    public_jobs = client.get(reverse("api-public-jobs", args=[organization.slug])).json()
    assert public_jobs["success"] is True
    assert public_jobs["data"]["results"] == []

    organization.status = Organization.Status.ACTIVE
    organization.save(update_fields=["status", "updated_at"])
    job.deleted_at = job.updated_at
    job.save(update_fields=["deleted_at", "updated_at"])
    assert client.get(reverse("public-job-detail", args=[organization.slug, job.slug])).status_code == 404
    assert client.get(reverse("api-public-job", args=[organization.slug, job.slug])).status_code == 404


@pytest.mark.django_db
def test_invalid_mutation_inputs_return_controlled_errors(
    client, organization, admin_user, application, job
):
    user, _ = admin_user
    client.force_login(user)
    move_response = client.post(
        reverse("api-move-application", args=[organization.slug, uuid.uuid4()]),
        data={"target_stage_id": str(application.current_stage_id), "expected_version": 1},
        content_type="application/json",
        HTTP_IDEMPOTENCY_KEY="missing-application",
    )
    assert move_response.status_code == 400
    assert move_response.json()["error"]["message"] == "Application not found."

    template = ScorecardTemplate.objects.create(job=job, name="Invalid rating", version=2)
    criterion = ScorecardCriterion.objects.create(
        template=template, competency="Communication", position=1
    )
    interview = Interview.objects.create(
        application=application,
        scorecard_template=template,
        template_snapshot={
            "name": template.name,
            "version": template.version,
            "criteria": [{"id": str(criterion.id), "competency": criterion.competency}],
        },
        title="Rating validation",
        start_at="2026-08-25T10:00:00Z",
        end_at="2026-08-25T11:00:00Z",
        created_by=user,
    )
    InterviewParticipant.objects.create(interview=interview, user=user)
    scorecard_response = client.post(
        reverse("scorecard", args=[organization.slug, interview.id]),
        {f"rating_{criterion.id}": "not-a-number", "action": "draft"},
    )
    assert scorecard_response.status_code == 200
    assert b"Ratings must be between 1 and 5" in scorecard_response.content

    mfa_response = client.post(reverse("mfa-setup"), {"token": "abcdef"})
    assert mfa_response.status_code == 200
    assert b"6 to 8 digit" in mfa_response.content


@pytest.mark.django_db
def test_duplicate_merge_handles_shared_skills_without_breaking_transaction(
    organization, admin_user, candidate
):
    user, membership = admin_user
    master = Candidate(organization=organization, first_name="Nadia", last_name="P.")
    master.set_email("nadia.other@example.test")
    master.save()
    skill = Skill.objects.create(organization=organization, name="Python")
    candidate.candidate_skills.create(skill=skill, source="RESUME")
    master.candidate_skills.create(skill=skill, source="RESUME")
    suggestion = DuplicateSuggestion.objects.create(
        organization=organization,
        candidate=candidate,
        possible_duplicate=master,
        score=95,
    )

    merged = merge_candidates(
        suggestion_id=suggestion.id,
        organization=organization,
        actor=user,
        membership=membership,
    )

    candidate.refresh_from_db()
    suggestion.refresh_from_db()
    assert merged == master
    assert candidate.deleted_at is not None
    assert master.candidate_skills.filter(skill=skill, source="RESUME").count() == 1
    assert suggestion.status == DuplicateSuggestion.Status.MERGED


@pytest.mark.django_db
def test_pipeline_exposes_and_processes_rejection_form(
    client, organization, admin_user, application, job
):
    user, _ = admin_user
    reason = RejectionReason.objects.create(
        organization=organization,
        name="Position filled",
        category="BUSINESS",
    )
    client.force_login(user)
    pipeline_page = client.get(reverse("pipeline-job", args=[organization.slug, job.id]))
    assert pipeline_page.status_code == 200
    assert b"Reject application" in pipeline_page.content
    assert b"Position filled" in pipeline_page.content

    response = client.post(
        reverse("reject-application", args=[organization.slug, application.id]),
        {"reason": str(reason.id), "note": "Role closed"},
    )
    assert response.status_code == 302
    application.refresh_from_db()
    assert application.status == Application.Status.REJECTED
    assert application.rejection_reason == reason
